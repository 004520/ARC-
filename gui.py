"""
人脸识别系统 — Tkinter GUI

运行方式：python gui.py

功能页面：
  - 主页：摄像头实时画面 + 识别/录入/脸库 按钮
  - 录入：摄像头预览 + 姓名输入 + 拍照录入 / 上传照片录入
  - 脸库：网格展示已注册人脸（照片、姓名、时间、删除）
"""

import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import cv2
import numpy as np
import time
import os
import sys
from PIL import Image, ImageDraw, ImageFont, ImageTk

if sys.platform == "win32":
    import ctypes
    try:
        ctypes.windll.shcore.SetProcessDpiAwareness(1)
    except Exception:
        pass
    try:
        import winsound
    except ImportError:
        winsound = None

from inference import FaceEngine
from database import FaceDatabase

WIN_W, WIN_H = 1200, 900
CAM_W, CAM_H = 800, 600
PHOTO_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "faces")
STRANGER_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "strangers")

# 浅色系配色
BG       = "#F0F2F5"
CARD     = "#FFFFFF"
PRIMARY  = "#4A90D9"
SUCCESS  = "#27AE60"
DANGER   = "#E74C3C"
PURPLE   = "#8E44AD"
WARNING  = "#FF9800"
TXT      = "#333333"
TXT2     = "#999999"
BORDER   = "#E0E0E0"

# PIL 中文字体
_font_cache = {}

def _load_font(size):
    if size in _font_cache:
        return _font_cache[size]
    # 按优先级搜索系统中文字体
    candidates = [
        "msyh.ttc", "msyhbd.ttc", "msyh.ttl", "msyhbd.ttl",  # 微软雅黑
        "simhei.ttf", "simhei.bold.ttf",                       # 黑体
        "simsun.ttc", "simsun.ttf",                             # 宋体
        "STZHONGS.TTF", "STKAITI.TTF",                         # 华文中宋/楷体
    ]
    for name in candidates:
        p = os.path.join("C:/Windows/Fonts", name)
        if os.path.isfile(p):
            _font_cache[size] = ImageFont.truetype(p, size)
            return _font_cache[size]
    # 尝试 Pillow 内置字体
    try:
        font = ImageFont.truetype("DejaVuSans.ttf", size)
    except Exception:
        font = ImageFont.load_default()
    _font_cache[size] = font
    return font

FONT_SMALL  = _load_font(18)
FONT_LARGE  = _load_font(36)


def _make_placeholder(text):
    img = Image.new("RGB", (120, 120), "#E0E0E0")
    draw = ImageDraw.Draw(img)
    bb = draw.textbbox((0, 0), text, font=FONT_SMALL)
    tw, th = bb[2] - bb[0], bb[3] - bb[1]
    draw.text(((120 - tw) // 2, (120 - th) // 2), text, fill="#AAA", font=FONT_SMALL)
    return img


class FaceAppGUI:
    def __init__(self):
        self.root = tk.Tk()
        self.root.title("人脸识别系统")
        self.root.geometry(f"{WIN_W}x{WIN_H}")
        self.root.resizable(False, False)
        self.root.configure(bg=BG)

        # 居中
        sw = self.root.winfo_screenwidth()
        sh = self.root.winfo_screenheight()
        self.root.geometry(f"{WIN_W}x{WIN_H}+{(sw-WIN_W)//2}+{(sh-WIN_H)//2}")

        # 模型 & 数据库（带异常捕获）
        try:
            self.engine = FaceEngine()
        except Exception as e:
            messagebox.showerror("初始化错误", f"模型加载失败：{e}\n请检查 checkpoints/ 目录下是否有模型文件")
            self.root.destroy()
            return
        try:
            self.db = FaceDatabase()
        except Exception as e:
            messagebox.showerror("初始化错误", f"数据库连接失败：{e}\n请检查 MySQL 服务是否启动")
            self.root.destroy()
            return

        # 摄像头
        self.cap = cv2.VideoCapture(0, cv2.CAP_DSHOW)
        self.cap.set(cv2.CAP_PROP_FRAME_WIDTH, CAM_W)
        self.cap.set(cv2.CAP_PROP_FRAME_HEIGHT, CAM_H)
        if not self.cap.isOpened():
            messagebox.showerror("错误", "摄像头打开失败，请检查设备")
            self.root.destroy()
            return

        self.cascade = cv2.CascadeClassifier(
            cv2.data.haarcascades + "haarcascade_frontalface_default.xml"
        )
        if self.cascade.empty():
            messagebox.showerror("错误", "人脸检测模型加载失败，请检查 OpenCV 数据目录")
            self.root.destroy()
            return
        self.frame = None
        self.faces = []
        self.recognized = False
        self.rec_results = []
        self.running = True
        self.page = "main"
        self._photo_ref = None
        self.attendance_mode = False
        self._flash_active = False

        os.makedirs(PHOTO_DIR, exist_ok=True)
        os.makedirs(STRANGER_DIR, exist_ok=True)

        self.container = tk.Frame(self.root, bg=BG)
        self.container.pack(fill="both", expand=True)
        self.pages = {}

        self._build_main_page()
        self._build_register_page()
        self._build_library_page()
        self._build_attendance_page()

        self._goto("main")
        self._tick_camera()

    def run(self):
        self.root.protocol("WM_DELETE_WINDOW", self._quit)
        self.root.mainloop()

    def _quit(self):
        self.running = False
        self._flash_active = False
        try:
            self.cap.release()
        except Exception:
            pass
        try:
            self.db.close()
        except Exception:
            pass
        self.root.quit()
        self.root.destroy()

    def _tick_camera(self):
        if not self.running:
            return
        ret, frame = self.cap.read()
        if ret:
            self.frame = frame
            if self.page in ("main", "register"):
                gray = cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY)
                gray = cv2.equalizeHist(gray)
                det = self.cascade.detectMultiScale(
                    gray, scaleFactor=1.15, minNeighbors=5, minSize=(80, 80)
                )
                self.faces = [(x, y, x+w, y+h, 1.0) for x, y, w, h in det]

                disp = self._overlay(frame.copy())
                rgb = cv2.cvtColor(disp, cv2.COLOR_BGR2RGB)
                img = Image.fromarray(rgb).resize((CAM_W, CAM_H))
                self._photo_ref = ImageTk.PhotoImage(img)

                if self.page == "main":
                    self._main_cam.config(image=self._photo_ref)
                else:
                    self._reg_cam.config(image=self._photo_ref)

        self.root.after(30, self._tick_camera)

    def _overlay(self, frame):
        if not self.faces:
            rgb = cv2.cvtColor(frame, cv2.COLOR_BGR2RGB)
            pil = Image.fromarray(rgb)
            draw_tmp = ImageDraw.Draw(pil)
            bb_check = draw_tmp.textbbox((0, 0), "未检测到人脸", font=FONT_LARGE)
            tw, th = bb_check[2] - bb_check[0], bb_check[3] - bb_check[1]
            cx, cy = pil.width // 2, pil.height // 2
            pad_x, pad_y = 30, 12
            rx1, ry1 = cx - tw // 2 - pad_x, cy - th // 2 - pad_y
            rx2, ry2 = cx + tw // 2 + pad_x, cy + th // 2 + pad_y
            overlay = Image.new("RGBA", pil.size, (0, 0, 0, 0))
            draw = ImageDraw.Draw(overlay)
            draw.rectangle([rx1, ry1, rx2, ry2], fill=(0, 0, 0, 140))
            pil = Image.alpha_composite(pil.convert("RGBA"), overlay).convert("RGB")
            draw = ImageDraw.Draw(pil)
            draw.text((cx - tw // 2, cy - th // 2), "未检测到人脸", fill=(255, 255, 255), font=FONT_LARGE)
            return cv2.cvtColor(np.array(pil), cv2.COLOR_RGB2BGR)

        rgb = cv2.cvtColor(frame, cv2.COLOR_BGR2RGB)
        pil = Image.fromarray(rgb)
        draw = ImageDraw.Draw(pil)
        img_h, img_w = pil.height, pil.width

        for i, (x1, y1, x2, y2, _) in enumerate(self.faces):
            is_register_page = (self.page == "register")
            if not is_register_page and self.recognized and self.rec_results and i < len(self.rec_results):
                name, sim = self.rec_results[i]
                is_known = (name != "陌生人")
                box_color = (0, 200, 0) if is_known else (220, 60, 60)
            else:
                box_color = (80, 140, 220)
                name, sim, is_known = None, None, False

            draw.rectangle([x1, y1, x2, y2], outline=box_color, width=3)

            if not is_register_page and self.recognized and self.rec_results and i < len(self.rec_results):
                if is_known:
                    label = f"{name}  {sim:.2f}"
                    label_bg = (0, 180, 0)
                else:
                    label = "陌生人"
                    label_bg = (220, 60, 60)

                bb = draw.textbbox((0, 0), label, font=FONT_SMALL)
                tw, th = bb[2] - bb[0], bb[3] - bb[1]
                lx = x1
                ly = y2 + 4
                if ly + th + 6 > img_h:
                    ly = y2 - th - 6
                if lx + tw + 10 > img_w:
                    lx = img_w - tw - 10
                if lx < 0:
                    lx = 0
                draw.rectangle([lx, ly, lx + tw + 10, ly + th + 6], fill=label_bg)
                draw.text((lx + 5, ly + 3), label, fill=(255, 255, 255), font=FONT_SMALL)

        return cv2.cvtColor(np.array(pil), cv2.COLOR_RGB2BGR)

    def _goto(self, name):
        self.page = name
        self.recognized = False
        self.rec_results = []
        self._flash_active = False
        for p in self.pages.values():
            p.pack_forget()
        self.pages[name].pack(fill="both", expand=True)
        if name == "library":
            self._refresh_library()
        elif name == "attendance":
            self._refresh_attendance()

    def _build_main_page(self):
        f = tk.Frame(self.container, bg=BG)
        self.pages["main"] = f

        tk.Label(f, text="人脸识别系统", font=("Microsoft YaHei", 20, "bold"),
                 bg=BG, fg=TXT).pack(pady=(15, 10))

        # 摄像头卡片
        card = tk.Frame(f, bg=CARD, highlightthickness=1, highlightbackground=BORDER)
        card.pack(padx=20, pady=5)
        self._main_cam = tk.Label(card, bg="black")
        self._main_cam.pack(padx=2, pady=2)

        # 状态栏
        self._main_status = tk.Label(f, text="就绪 — 请点击下方按钮", font=("Microsoft YaHei", 11),
                                     bg=BG, fg=TXT2)
        self._main_status.pack(pady=5)

        # 统计面板
        stat_card = tk.Frame(f, bg=CARD, highlightthickness=1, highlightbackground=BORDER)
        stat_card.pack(padx=20, pady=(0, 5), fill="x")
        self._stat_label = tk.Label(
            stat_card, text="加载中...", font=("Microsoft YaHei", 10),
            bg=CARD, fg=TXT2, pady=6,
        )
        self._stat_label.pack()
        self._tick_stats()

        # 按钮行
        bf = tk.Frame(f, bg=BG)
        bf.pack(pady=10)
        self._btn(bf, "识  别", PRIMARY, self._do_recognize).pack(side="left", padx=12)
        self._btn(bf, "录  入", SUCCESS, lambda: self._goto("register")).pack(side="left", padx=12)
        self._btn(bf, "脸  库", PURPLE, lambda: self._goto("library")).pack(side="left", padx=12)

        # 考勤按钮行
        af = tk.Frame(f, bg=BG)
        af.pack(pady=(0, 5))
        self._att_btn = tk.Button(
            af, text="考勤模式：关", font=("Microsoft YaHei", 9),
            bg=TXT2, fg="white", bd=0, cursor="hand2",
            width=14, height=2,
            command=self._toggle_attendance,
        )
        self._att_btn.pack(side="left", padx=12)
        self._btn(af, "考勤记录", WARNING, lambda: self._goto("attendance")).pack(
            side="left", padx=12
        )

    def _build_register_page(self):
        f = tk.Frame(self.container, bg=BG)
        self.pages["register"] = f

        hdr = tk.Frame(f, bg=BG)
        hdr.pack(fill="x", padx=20, pady=(10, 5))
        self._btn(hdr, "< 返回", TXT2, lambda: self._goto("main"), w=8, small=True).pack(side="left")
        tk.Label(hdr, text="录入人脸", font=("Microsoft YaHei", 18, "bold"),
                 bg=BG, fg=TXT).pack(side="left", padx=20)

        card = tk.Frame(f, bg=CARD, highlightthickness=1, highlightbackground=BORDER)
        card.pack(padx=20, pady=5)
        self._reg_cam = tk.Label(card, bg="black")
        self._reg_cam.pack(padx=2, pady=2)

        # 姓名输入
        inp = tk.Frame(f, bg=CARD, highlightthickness=1, highlightbackground=BORDER)
        inp.pack(padx=20, pady=8, fill="x")
        row = tk.Frame(inp, bg=CARD)
        row.pack(pady=12, padx=20)
        tk.Label(row, text="姓名：", font=("Microsoft YaHei", 12), bg=CARD, fg=TXT).pack(side="left")
        self._reg_name = tk.StringVar()
        e = tk.Entry(row, textvariable=self._reg_name, font=("Microsoft YaHei", 12),
                     width=18, bd=1, relief="solid")
        e.pack(side="left", padx=8)
        e.focus_set()

        # 按钮
        bf = tk.Frame(f, bg=BG)
        bf.pack(pady=6)
        self._btn(bf, "录入人脸", SUCCESS, self._do_reg_cam).pack(side="left", padx=12)
        self._btn(bf, "上传照片录入", PRIMARY, self._do_reg_photo).pack(side="left", padx=12)

        self._reg_status = tk.Label(f, text="", font=("Microsoft YaHei", 11), bg=BG, fg=TXT2)
        self._reg_status.pack(pady=4)

    def _build_library_page(self):
        f = tk.Frame(self.container, bg=BG)
        self.pages["library"] = f

        hdr = tk.Frame(f, bg=BG)
        hdr.pack(fill="x", padx=20, pady=(10, 5))
        self._btn(hdr, "< 返回", TXT2, lambda: self._goto("main"), w=8, small=True).pack(side="left")
        self._lib_title = tk.Label(hdr, text="人脸库", font=("Microsoft YaHei", 18, "bold"),
                                   bg=BG, fg=TXT)
        self._lib_title.pack(side="left", padx=20)

        # 可滚动区域
        canvas = tk.Canvas(f, bg=BG, highlightthickness=0)
        vsb = ttk.Scrollbar(f, orient="vertical", command=canvas.yview)
        self._lib_inner = tk.Frame(canvas, bg=BG)
        self._lib_inner.bind("<Configure>",
                             lambda e: canvas.configure(scrollregion=canvas.bbox("all")))
        canvas.create_window((0, 0), window=self._lib_inner, anchor="nw")
        canvas.configure(yscrollcommand=vsb.set)
        canvas.pack(side="left", fill="both", expand=True, padx=(20, 0), pady=10)
        vsb.pack(side="right", fill="y", padx=(0, 20), pady=10)

        # 鼠标滚轮
        self._lib_canvas = canvas
        canvas.bind("<Enter>", lambda e: canvas.bind_all("<MouseWheel>",
                    lambda ev: canvas.yview_scroll(int(-ev.delta / 120), "units")))
        canvas.bind("<Leave>", lambda e: canvas.unbind_all("<MouseWheel>"))

    def _refresh_library(self):
        for w in self._lib_inner.winfo_children():
            w.destroy()

        users = self.db.list_users()
        self._lib_title.config(text=f"人脸库（{len(users)} 人）")

        if not users:
            tk.Label(self._lib_inner, text="暂无录入数据", font=("Microsoft YaHei", 14),
                     bg=BG, fg=TXT2).pack(pady=60)
            return

        cols = 4
        for i, u in enumerate(users):
            self._lib_card(self._lib_inner, u).grid(
                row=i // cols, column=i % cols, padx=8, pady=8, sticky="nsew")
        for c in range(cols):
            self._lib_inner.columnconfigure(c, weight=1)

    def _lib_card(self, parent, user):
        card = tk.Frame(parent, bg=CARD, highlightthickness=1, highlightbackground=BORDER)

        # 头像
        pp = user.get("photo_path", "")
        if pp and os.path.isfile(pp):
            try:
                thumb = Image.open(pp).convert("RGB").resize((120, 120))
                photo = ImageTk.PhotoImage(thumb)
            except Exception:
                ph = _make_placeholder("无照片")
                photo = ImageTk.PhotoImage(ph)
        else:
            ph = _make_placeholder("无照片")
            photo = ImageTk.PhotoImage(ph)

        lbl = tk.Label(card, image=photo, bg=CARD, width=120, height=120)
        lbl.image = photo
        lbl.pack(pady=(10, 4))

        # 姓名
        tk.Label(card, text=user["name"], font=("Microsoft YaHei", 11, "bold"),
                 bg=CARD, fg=TXT).pack()

        # 时间
        ct = str(user.get("created_at", ""))
        if ct:
            ct = ct.split(".")[0]
        tk.Label(card, text=ct, font=("Microsoft YaHei", 9), bg=CARD, fg=TXT2).pack()

        # 删除
        uid = user["user_id"]
        uname = user["name"]
        self._btn(card, "删除", DANGER, lambda u=uid, n=uname: self._del_user(u, n),
                  w=7, small=True).pack(pady=(4, 10))
        return card

    def _del_user(self, uid, name):
        if messagebox.askyesno("确认删除", f"确定删除「{name}」的人脸数据吗？"):
            pp = os.path.join(PHOTO_DIR, f"{uid}.jpg")
            if os.path.exists(pp):
                os.remove(pp)
            self.db.delete_user(uid)
            self._refresh_library()

    def _toggle_attendance(self):
        self.attendance_mode = not self.attendance_mode
        if self.attendance_mode:
            self._att_btn.config(text="考勤模式：开", bg=SUCCESS)
            self._main_status.config(text="考勤模式已开启 — 识别成功后自动打卡", fg=SUCCESS)
        else:
            self._att_btn.config(text="考勤模式：关", bg=TXT2)
            self._main_status.config(text="考勤模式已关闭", fg=TXT2)

    def _build_attendance_page(self):
        f = tk.Frame(self.container, bg=BG)
        self.pages["attendance"] = f

        # 顶部栏
        hdr = tk.Frame(f, bg=BG)
        hdr.pack(fill="x", padx=20, pady=(10, 5))
        self._btn(hdr, "< 返回", TXT2, lambda: self._goto("main"), w=8, small=True).pack(side="left")
        tk.Label(hdr, text="考勤记录", font=("Microsoft YaHei", 18, "bold"),
                 bg=BG, fg=TXT).pack(side="left", padx=20)

        # 筛选栏
        filt = tk.Frame(f, bg=CARD, highlightthickness=1, highlightbackground=BORDER)
        filt.pack(fill="x", padx=20, pady=(8, 0))

        tk.Label(filt, text="起始日期：", font=("Microsoft YaHei", 10), bg=CARD, fg=TXT).pack(
            side="left", padx=(15, 2))
        self._att_from = tk.StringVar(value="")
        tk.Entry(filt, textvariable=self._att_from, font=("Microsoft YaHei", 10),
                 width=12, bd=1, relief="solid").pack(side="left", padx=2)

        tk.Label(filt, text="截止日期：", font=("Microsoft YaHei", 10), bg=CARD, fg=TXT).pack(
            side="left", padx=(10, 2))
        self._att_to = tk.StringVar(value="")
        tk.Entry(filt, textvariable=self._att_to, font=("Microsoft YaHei", 10),
                 width=12, bd=1, relief="solid").pack(side="left", padx=2)

        self._btn(filt, "查询", PRIMARY, self._refresh_attendance, w=8, small=True).pack(
            side="left", padx=15, pady=8)

        self._btn(filt, "导出 Excel", WARNING, self._export_attendance, w=10, small=True).pack(
            side="right", padx=15, pady=8)

        # 统计栏
        self._att_stat = tk.Label(f, text="", font=("Microsoft YaHei", 10), bg=BG, fg=TXT2)
        self._att_stat.pack(pady=5)

        # TreeView 表格
        tree_frame = tk.Frame(f, bg=BG)
        tree_frame.pack(fill="both", expand=True, padx=20, pady=(0, 15))

        columns = ("name", "date", "time", "similarity", "status")
        self._att_tree = ttk.Treeview(
            tree_frame, columns=columns, show="headings",
            height=20, selectmode="browse",
        )
        self._att_tree.heading("name", text="姓名")
        self._att_tree.heading("date", text="日期")
        self._att_tree.heading("time", text="首次打卡")
        self._att_tree.heading("similarity", text="识别次数")
        self._att_tree.heading("status", text="状态")

        self._att_tree.column("name", width=120, anchor="center")
        self._att_tree.column("date", width=120, anchor="center")
        self._att_tree.column("time", width=120, anchor="center")
        self._att_tree.column("similarity", width=100, anchor="center")
        self._att_tree.column("status", width=80, anchor="center")

        vsb = ttk.Scrollbar(tree_frame, orient="vertical", command=self._att_tree.yview)
        self._att_tree.configure(yscrollcommand=vsb.set)
        self._att_tree.pack(side="left", fill="both", expand=True)
        vsb.pack(side="right", fill="y")

        style = ttk.Style()
        style.configure("Treeview", font=("Microsoft YaHei", 10), rowheight=28)
        style.configure("Treeview.Heading", font=("Microsoft YaHei", 10, "bold"))

    def _refresh_attendance(self):
        """从数据库加载考勤记录并填充表格。"""
        for row in self._att_tree.get_children():
            self._att_tree.delete(row)

        date_from = self._att_from.get().strip() or None
        date_to = self._att_to.get().strip() or None

        records = self.db.get_attendance_stats(date_from=date_from, date_to=date_to)

        if not records:
            self._att_stat.config(text="暂无考勤记录")
            return

        late_count = 0
        for r in records:
            name = r["name"]
            att_date = str(r["att_date"])
            first_time = str(r["first_time"])
            count = r["rec_count"]

            status = "迟到" if first_time > "09:00:00" else "正常"
            if status == "迟到":
                late_count += 1
                tag = "late"
            else:
                tag = "normal"

            self._att_tree.insert("", "end", values=(
                name, att_date, first_time,
                f"{count}次", status,
            ), tags=(tag,))

        self._att_tree.tag_configure("normal", foreground=SUCCESS)
        self._att_tree.tag_configure("late", foreground=DANGER)

        total_people = len(set(r["name"] for r in records))
        total_days = len(set(str(r["att_date"]) for r in records))
        self._att_stat.config(
            text=f"共 {total_people} 人，{total_days} 天，{len(records)} 条记录  |  "
                 f"迟到 {late_count} 次  |  默认筛选最近 7 天"
        )

    def _export_attendance(self):
        """导出考勤记录到 Excel，用户选择保存路径。"""
        date_from = self._att_from.get().strip() or None
        date_to = self._att_to.get().strip() or None

        records = self.db.get_attendance_stats(date_from=date_from, date_to=date_to)
        if not records:
            messagebox.showinfo("提示", "没有可导出的考勤记录")
            return

        path = filedialog.asksaveasfilename(
            title="导出考勤报表",
            defaultextension=".xlsx",
            filetypes=[
                ("Excel 文件", "*.xlsx"),
                ("CSV 文件", "*.csv"),
            ],
            initialfile=f"考勤报表_{time.strftime('%Y%m%d')}.xlsx",
        )
        if not path:
            return

        try:
            import openpyxl
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "考勤记录"

            headers = ["姓名", "日期", "首次打卡时间", "识别次数", "状态"]
            header_fill = openpyxl.styles.PatternFill(
                start_color="4A90D9", end_color="4A90D9", fill_type="solid"
            )
            header_font = openpyxl.styles.Font(bold=True, color="FFFFFF", size=12)
            thin_border = openpyxl.styles.Border(
                left=openpyxl.styles.Side(style="thin"),
                right=openpyxl.styles.Side(style="thin"),
                top=openpyxl.styles.Side(style="thin"),
                bottom=openpyxl.styles.Side(style="thin"),
            )
            center_align = openpyxl.styles.Alignment(horizontal="center", vertical="center")

            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.fill = header_fill
                cell.font = header_font
                cell.border = thin_border
                cell.alignment = center_align

            for row_idx, r in enumerate(records, 2):
                first_time = str(r["first_time"])
                status = "迟到" if first_time > "09:00:00" else "正常"
                values = [
                    r["name"],
                    str(r["att_date"]),
                    first_time,
                    r["rec_count"],
                    status,
                ]
                for col, val in enumerate(values, 1):
                    cell = ws.cell(row=row_idx, column=col, value=val)
                    cell.border = thin_border
                    cell.alignment = center_align
                    if status == "迟到" and col == 5:
                        cell.font = openpyxl.styles.Font(color="E74C3C", bold=True)

            for col in range(1, 6):
                ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 18

            wb.save(path)
            messagebox.showinfo("导出成功", f"考勤报表已保存到：\n{path}")

        except ImportError:
            import csv
            with open(path, "w", newline="", encoding="utf-8-sig") as f:
                writer = csv.writer(f)
                writer.writerow(["姓名", "日期", "首次打卡时间", "识别次数", "状态"])
                for r in records:
                    first_time = str(r["first_time"])
                    status = "迟到" if first_time > "09:00:00" else "正常"
                    writer.writerow([
                        r["name"], str(r["att_date"]), first_time,
                        r["rec_count"], status,
                    ])
            messagebox.showinfo("导出成功", f"考勤报表已保存到：\n{path}")

        except Exception as e:
            messagebox.showerror("导出失败", f"导出时发生错误：{e}")

    def _do_recognize(self):
        """点击识别按钮 → 识别当前所有检测到的人脸"""
        if not self.faces:
            self._main_status.config(text="未检测到人脸", fg=DANGER)
            return

        gf, gids, gnames = self.db.load_all_features_with_names()
        if len(gids) == 0:
            self._main_status.config(text="脸库为空，请先录入人脸", fg=DANGER)
            return

        self._main_status.config(text="识别中...", fg=PRIMARY)
        self.root.update_idletasks()

        if self.frame is None:
            self._main_status.config(text="摄像头画面未就绪", fg=DANGER)
            return
        h, w = self.frame.shape[:2]
        results = []
        for x1, y1, x2, y2, _ in self.faces:
            m = 0.3
            mx, my = int((x2 - x1) * m), int((y2 - y1) * m)
            crop = self.frame[max(0, y1 - my):min(h, y2 + my),
                              max(0, x1 - mx):min(w, x2 + mx)]
            if crop.size == 0:
                results.append(("陌生人", 0.0))
                continue
            feat = self.engine.extract_array(crop)
            sims = gf @ feat
            idx = int(np.argmax(sims))
            ms = float(sims[idx])
            if ms >= 0.45:
                results.append((gnames[idx], ms))
            else:
                results.append(("陌生人", ms))

        self.recognized = True
        self.rec_results = results
        self._flash_active = False

        if self.attendance_mode:
            for name, sim in results:
                if name != "陌生人":
                    try:
                        self.db.log_recognition(
                            user_id="", name=name, similarity=sim,
                        )
                    except Exception:
                        pass

        names = [r[0] for r in results]
        known_names = [n for n in names if n != "陌生人"]
        has_stranger = any(n == "陌生人" for n in names)
        if known_names:
            self._main_status.config(
                text=f"识别成功：{'、'.join(known_names)}", fg=SUCCESS
            )
        elif has_stranger:
            self._main_status.config(
                text="⚠ 警告：陌生人！未匹配到已知人脸", fg=DANGER
            )
        else:
            self._main_status.config(text="识别失败：未匹配到已知人脸", fg=DANGER)

        if has_stranger:
            self._save_stranger_snapshot()
            self._stranger_alert()

    def _save_stranger_snapshot(self):
        """陌生人自动截图保存到 strangers/ 目录"""
        if self.frame is None:
            return
        timestamp = time.strftime("%Y%m%d_%H%M%S")
        filepath = os.path.join(STRANGER_DIR, f"stranger_{timestamp}.jpg")
        cv2.imwrite(filepath, self.frame)

    def _stranger_alert(self):
        if sys.platform == "win32" and winsound:
            try:
                winsound.Beep(1000, 300)
            except Exception:
                pass
        self._flash_active = True
        self._flash_count = 6
        self._flash_alert()

    def _flash_alert(self):
        if not self._flash_active or self._flash_count <= 0:
            self._flash_active = False
            self._main_status.config(
                text="⚠ 警告：陌生人！未匹配到已知人脸", fg=DANGER
            )
            return
        if self._flash_count % 2 == 0:
            self._main_status.config(fg=DANGER)
        else:
            self._main_status.config(fg="white")
        self._flash_count -= 1
        self.root.after(200, self._flash_alert)

    def _tick_stats(self):
        """定时刷新主页统计面板"""
        if not self.running:
            return
        try:
            total_users = self.db.count()
            total_logs = self.db.count_logs()
            today_str = time.strftime("%Y-%m-%d")
            self.db.cursor.execute(
                "SELECT COUNT(*) FROM recognition_logs WHERE DATE(created_at) = %s",
                (today_str,),
            )
            today_logs = self.db.cursor.fetchone()[0]
            self.db.cursor.execute(
                "SELECT COUNT(*) FROM recognition_logs WHERE DATE(created_at) = %s AND name = '陌生人'",
                (today_str,),
            )
            today_strangers = self.db.cursor.fetchone()[0]
            self.db.cursor.execute(
                "SELECT name, created_at FROM recognition_logs ORDER BY created_at DESC LIMIT 1"
            )
            last_row = self.db.cursor.fetchone()
            last_info = f"{last_row[0]}（{str(last_row[1]).split('.')[0]}）" if last_row else "无"

            self._stat_label.config(
                text=f"脸库: {total_users}人  |  今日识别: {today_logs}次  |  "
                     f"今日陌生人: {today_strangers}次  |  最后识别: {last_info}  |  "
                     f"累计识别: {total_logs}次",
            )
        except Exception:
            self._stat_label.config(text="统计信息加载失败")
        self.root.after(30000, self._tick_stats)

    def _do_reg_cam(self):
        """从摄像头拍照录入"""
        name = self._reg_name.get().strip()
        if not name:
            self._reg_status.config(text="请先输入姓名", fg=DANGER)
            return
        if not self.faces:
            self._reg_status.config(text="未检测到人脸，请面对摄像头", fg=DANGER)
            return

        if self.frame is None:
            self._reg_status.config(text="摄像头画面未就绪", fg=DANGER)
            return
        h, w = self.frame.shape[:2]
        # 取最大人脸
        best = max(self.faces, key=lambda f: (f[2] - f[0]) * (f[3] - f[1]))
        x1, y1, x2, y2, _ = best

        quality_msg = self._check_face_quality(x1, y1, x2, y2, w, h)
        if quality_msg:
            self._reg_status.config(text=quality_msg, fg=WARNING)
            return

        m = 0.3
        mx, my = int((x2 - x1) * m), int((y2 - y1) * m)
        crop = self.frame[max(0, y1 - my):min(h, y2 + my),
                          max(0, x1 - mx):min(w, x2 + mx)]
        if crop.size == 0:
            self._reg_status.config(text="人脸裁剪失败", fg=DANGER)
            return

        feat = self.engine.extract_array(crop)
        uid = f"user_{int(time.time())}"

        pp = os.path.join(PHOTO_DIR, f"{uid}.jpg")
        crop_rgb = crop[:, :, ::-1] if crop.ndim == 3 and crop.shape[2] == 3 else crop
        Image.fromarray(crop_rgb.astype(np.uint8), "RGB").save(pp, quality=95)

        self.db.register_with_photo(uid, name, feat, pp)

        total = self.db.count()
        self._reg_status.config(
            text=f"录入成功：{name}（脸库共 {total} 人） | 人脸质量：优秀",
            fg=SUCCESS,
        )
        self._reg_name.set("")

    def _check_face_quality(self, x1, y1, x2, y2, img_w, img_h):
        """检测人脸质量，返回提示字符串或 None（通过）。"""
        face_w = x2 - x1
        face_h = y2 - y1
        face_area = face_w * face_h
        img_area = img_w * img_h

        if face_area < img_area * 0.03:
            return "人脸过小，请靠近摄像头"

        if face_area > img_area * 0.80:
            return "人脸过大，请远离摄像头"

        face_cx = (x1 + x2) / 2
        face_cy = (y1 + y2) / 2
        if (abs(face_cx - img_w / 2) > img_w * 0.30 or
                abs(face_cy - img_h / 2) > img_h * 0.30):
            return "请将人脸移至画面中央"

        if self.frame is not None:
            gray_face = cv2.cvtColor(
                self.frame[max(0, y1):min(self.frame.shape[0], y2),
                           max(0, x1):min(self.frame.shape[1], x2)],
                cv2.COLOR_BGR2GRAY,
            )
            if gray_face.size > 0 and np.mean(gray_face) < 50:
                return "光线偏暗，请调整光照条件"

        return None

    def _do_reg_photo(self):
        """上传照片录入"""
        name = self._reg_name.get().strip()
        if not name:
            self._reg_status.config(text="请先输入姓名", fg=DANGER)
            return

        path = filedialog.askopenfilename(title="选择照片",
                                          filetypes=[("图片", "*.jpg *.jpeg *.png *.bmp")])
        if not path:
            return

        img = cv2.imread(path, cv2.IMREAD_COLOR)
        if img is None:
            self._reg_status.config(text="图片读取失败", fg=DANGER)
            return

        gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
        gray = cv2.equalizeHist(gray)
        det = self.cascade.detectMultiScale(gray, scaleFactor=1.15, minNeighbors=5, minSize=(80, 80))
        if len(det) == 0:
            self._reg_status.config(text="未在照片中检测到人脸", fg=DANGER)
            return

        hi, wi = img.shape[:2]
        # 取最大人脸
        areas = [w * h for (x, y, w, h) in det]
        idx = int(np.argmax(areas))
        x, y, fw, fh = det[idx]

        face_area = fw * fh
        img_area = wi * hi
        if face_area < img_area * 0.03:
            self._reg_status.config(text="照片中人脸过小，请选择更清晰的照片", fg=WARNING)
            return
        gray_face = cv2.cvtColor(img[max(0, y):min(hi, y+fh), max(0, x):min(wi, x+fw)], cv2.COLOR_BGR2GRAY)
        if gray_face.size > 0 and np.mean(gray_face) < 50:
            self._reg_status.config(text="照片光线偏暗，请选择更明亮清晰的照片", fg=WARNING)
            return

        m = 0.3
        mx, my = int(fw * m), int(fh * m)
        crop = img[max(0, y - my):min(hi, y + fh + my),
                   max(0, x - mx):min(wi, x + fw + mx)]

        feat = self.engine.extract_array(crop)
        uid = f"user_{int(time.time())}"

        pp = os.path.join(PHOTO_DIR, f"{uid}.jpg")
        crop_rgb = crop[:, :, ::-1] if crop.ndim == 3 and crop.shape[2] == 3 else crop
        Image.fromarray(crop_rgb.astype(np.uint8), "RGB").save(pp, quality=95)

        self.db.register_with_photo(uid, name, feat, pp)

        total = self.db.count()
        self._reg_status.config(
            text=f"录入成功：{name}（脸库共 {total} 人） | 人脸质量：优秀",
            fg=SUCCESS,
        )
        self._reg_name.set("")

    def _btn(self, parent, text, color, cmd, w=12, small=False):
        sz = 9 if small else 12
        h = 1 if small else 2
        b = tk.Button(parent, text=text, font=("Microsoft YaHei", sz),
                      bg=color, fg="white", activebackground=color, activeforeground="white",
                      bd=0, cursor="hand2", width=w, height=h, command=cmd)
        b.bind("<Enter>", lambda e: b.config(bg=_darken(color)))
        b.bind("<Leave>", lambda e: b.config(bg=color))
        return b


def _darken(c, f=0.85):
    c = c.lstrip("#")
    r, g, b = (int(c[i:i+2], 16) for i in (0, 2, 4))
    return f"#{int(r*f):02x}{int(g*f):02x}{int(b*f):02x}"

if __name__ == "__main__":
    app = FaceAppGUI()
    if app.root.winfo_exists():
        app.run()
